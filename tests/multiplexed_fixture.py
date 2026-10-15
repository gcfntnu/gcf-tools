"""Private copies of synthetic paired inputs for multiplexed metadata checks."""
from pathlib import Path
import shutil

from openpyxl import load_workbook

RUN = "990101_SYNTHETIC_0001_ASYNTHETIC"
PROJECT = "GCF-2099-001"
KIT = "Parse Biosciences Test"
CUSTOMER = "Sample-Submission-Form"
LAB = "INFO (GCF-lab only)"
DEMUX = "Cell Multiplexing"


def make_multiplexed(tmp_path, sheet_kit=KIT):
    source = tmp_path / "input"
    shutil.copytree(Path(__file__).parent / "fixtures/paired", source)
    run = source / RUN
    sheet = run / "SampleSheet.csv"
    sheet.write_text(sheet.read_text().replace("Libprep,Synthetic", "Libprep," + sheet_kit))
    form = run / "Sample-Submission-Form.xlsx"
    book = load_workbook(form)
    customer = book[CUSTOMER]
    # Blank biological template columns on the library rows.
    for row in range(16, 19):
        customer.cell(row, 2).value = None
        customer.cell(row, 3).value = None
    for col, header in enumerate(["Library External ID", "Custom Metric [dtype=numerical]",
                                   "Volume", "Sample buffer", "Sample type", "Quantification"], 6):
        customer.cell(15, col, header)
    for row in range(16, 19):
        for col, value in enumerate(["LIB" + str(row), 0, 5, "buffer", "RNA", "method"], 6):
            customer.cell(row, col, value)
    lab = book[LAB]
    lab.cell(1, 3, "KIT")
    lab.cell(1, 4, "Custom Score [dtype=numerical]")
    for row in range(2, 5):
        lab.cell(row, 3, "synthetic")
        lab.cell(row, 4, 0)
    demux = book.create_sheet(DEMUX)
    demux.append(["Unique Sample ID", "Wells", "Sample Group", "Sample biosource", "External ID", "Custom Metric"])
    demux.append(["001", "A1, A2", "control", "sourceA", "007", 0])
    demux.append(["1", "B1-B2", "treated", "sourceB", "008", 2])
    book.save(form)
    work = tmp_path / "analysis"
    shutil.copytree(source / "workflow", work / "src/gcf-workflows")
    kit = work / "src/gcf-workflows/libprep.config"
    kit.write_text(kit.read_text().replace("Synthetic PE", KIT + " PE"))
    return run, form, work
