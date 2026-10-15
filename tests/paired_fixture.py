"""Generate the wholly synthetic #62 fixture; never imports the testdata producer."""
import argparse
from datetime import datetime
import gzip
import hashlib
import io
import json
from pathlib import Path
import re
import zipfile

from openpyxl import Workbook

PROJECT = "GCF-2099-001"
RUN = "990101_SYNTHETIC_0001_ASYNTHETIC"
SAMPLES = ("001", "001A", "CTRL")


def create_fixture(destination):
    """Write a new directory only; deterministic under requirements-dev.txt."""
    root = Path(destination)
    root.mkdir(parents=True, exist_ok=False)
    run = root / RUN
    project = run / PROJECT
    project.mkdir(parents=True)
    (run / "SampleSheet.csv").write_text(
        "[Header]\nIEMFileVersion,4\nExperimentName,Synthetic paired example\n"
        "[CustomOptions]\nLibprep,Synthetic\nOrganism,Homo sapiens\n"
        "[Data]\nSample_ID,Sample_Project\n" +
        "".join(f"{sample},{PROJECT}\n" for sample in SAMPLES)
    )
    book = Workbook()
    book.properties.creator = "gcf-tools synthetic fixture"
    book.properties.lastModifiedBy = "gcf-tools synthetic fixture"
    book.properties.created = book.properties.modified = datetime(2000, 1, 1)
    customer = book.active
    customer.title = "Sample-Submission-Form"
    headers = ["Unique Sample ID", "External ID", "Sample Group", "Project ID", "Concentration"]
    for column, value in enumerate(headers, 1):
        customer.cell(15, column, value)
    for row, values in enumerate([
        ["001", "007", "control", PROJECT, 10],
        ["001A", "008", "treated", PROJECT, 20],
        ["CTRL", "009", "control", PROJECT, 30],
    ], 16):
        for column, value in enumerate(values, 1):
            customer.cell(row, column, value)
    lab = book.create_sheet("INFO (GCF-lab only)")
    lab.append(["Sample_ID", "Concentration"])
    for sample, concentration in zip(SAMPLES, [11, 22, 33]):
        lab.append([sample, concentration])
    buffer = io.BytesIO()
    book.save(buffer)
    # openpyxl sets ZIP times and modified date on save. Normalize both.
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(
            run / "Sample-Submission-Form.xlsx", "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            data = source.read(name)
            if name == "docProps/core.xml":
                data = re.sub(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", b"2000-01-01T00:00:00Z", data)
            info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, data)
    (run / "Stats").mkdir()
    (run / "Stats/Stats.json").write_text(json.dumps({
        "ReadInfosForLanes": [{"LaneNumber": 1, "ReadInfos": [
            {"Number": 1, "IsIndexedRead": False, "NumCycles": 75},
            {"Number": 2, "IsIndexedRead": False, "NumCycles": 150},
        ]}],
    }, indent=2) + "\n")
    checksums = []
    for sample in SAMPLES:
        for mate, length in [(1, 75), (2, 150)]:
            records = []
            for number in range(1, 25):
                # Distinct repeatable sequences, no PRNG/version-dependent seed.
                digest = hashlib.shake_256(f"{sample}:{number}:{mate}".encode()).digest(length)
                sequence = "".join("ACGT"[value % 4] for value in digest)
                records.append(f"@synthetic:{sample}:{number:04d}/{mate}\n{sequence}\n+\n{'I' * length}\n")
            path = project / f"{sample}_R{mate}.fastq.gz"
            with path.open("wb") as handle, gzip.GzipFile(filename="", mode="wb", fileobj=handle, mtime=0) as compressed:
                compressed.write("".join(records).encode())
            checksums.append(f"{hashlib.md5(path.read_bytes()).hexdigest()}  {PROJECT}/{path.name}\n")
    (run / f"md5sum_{PROJECT}_fastq.txt").write_text("".join(checksums))
    workflow = root / "workflow"
    (workflow / "rnaseq").mkdir(parents=True)
    (workflow / "libprep.config").write_text(
        "# Synthetic initialization-only stub, not a scientific kit definition.\n"
        "Synthetic PE:\n  workflow: rnaseq\n  library_layout: paired\n"
        "  filter:\n    trim:\n      fastp:\n        params: '-q 17'\n"
    )
    (workflow / "rnaseq/rnaseq.smk").write_text(
        '# Initialization-only stub; no scientific workflow is supplied.\n'
        'raise RuntimeError("Synthetic fixture is not an executable workflow")\n'
    )
    provenance = {
        "origin": "Wholly synthetic; no facility data or upstream workflow files copied",
        "generator": "tests/paired_fixture.py", "pairs_per_sample": 24,
        "geometry": [75, 150], "gzip": "mtime=0, no filename, compression level 9",
        "workflow": "Local initialization-only stub: workflow/libprep.config and workflow/rnaseq/rnaseq.smk",
        "optional_workflow_comparator": "42d37e302a2ff6c868844e6d33147460e0ec7ba3",
        "sha256": {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest()
                   for p in sorted(root.rglob("*")) if p.is_file()},
    }
    (root / "provenance.json").write_text(json.dumps(provenance, indent=2, sort_keys=True) + "\n")
    return root


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("destination", type=Path, help="New directory; existing paths are refused")
    create_fixture(parser.parse_args().destination)
