"""Safe SE stand-in for the Git-only legacy workbook, not the #62 PE fixture.

Created independently for packaging checks; no facility records or producer output.
The J48 note preserves the private-copy repair assertion in test_libprep_cli.
"""
import gzip
import json
from pathlib import Path

from openpyxl import Workbook


def create_run(destination):
    run = Path(destination)
    project = run / "GCF-2026-001"
    project.mkdir(parents=True)
    (run / "SampleSheet.csv").write_text(
        "[Header]\nIEMFileVersion,4\n[Data]\n"
        "Sample_ID,Sample_Project\n001,GCF-2026-001\n"
    )
    workbook = Workbook()
    customer = workbook.active
    customer.title = "Sample-Submission-Form"
    for column, value in enumerate(["Unique Sample ID", "External ID", "Project ID"], 1):
        customer.cell(15, column, value)
    for column, value in enumerate(["001", "synthetic", "GCF-2026-001"], 1):
        customer.cell(16, column, value)
    customer["J48"] = "Synthetic invalid note; repair private copy before CLI check"
    workbook.create_sheet("INFO (GCF-lab only)").append(["Sample_ID", "Concentration"])
    workbook.save(run / "Sample-Submission-Form.xlsx")
    (run / "Stats").mkdir()
    (run / "Stats/Stats.json").write_text(json.dumps({
        "ReadInfosForLanes": [{"LaneNumber": 1, "ReadInfos": [
            {"Number": 1, "IsIndexedRead": False, "NumCycles": 86},
        ]}],
    }))
    record = b"@synthetic\n" + b"A" * 86 + b"\n+\n" + b"I" * 86 + b"\n"
    (project / "001_R1.fastq.gz").write_bytes(gzip.compress(record, mtime=0))
    return run
