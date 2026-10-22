"""Synthetic installed-command regressions for analysis/source identity (#41)."""
import csv
import gzip
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys

import pytest
import yaml
from openpyxl import load_workbook

from test_paired_cli import FIXTURE, PROJECT, RUN, contents

SECOND = "GCF-2099-002"
ANALYSIS = "GCF-2099-099"


def make_sources(tmp_path, assignments):
    """Each run maps exact sample IDs to projects; repeats across runs are intentional."""
    root = tmp_path / "sources"
    runs = []
    for number, selected in enumerate(assignments):
        run = root / f"99010{number + 1}_SYNTHETIC_000{number + 1}_ABATCH{number + 1}"
        shutil.copytree(FIXTURE / RUN, run)
        shutil.rmtree(run / PROJECT)
        (run / f"md5sum_{PROJECT}_fastq.txt").unlink()
        (run / "SampleSheet.csv").write_text(
            "[Header]\nIEMFileVersion,4\n[CustomOptions]\nLibprep,Synthetic\n"
            "Organism,Homo sapiens\n[Data]\nSample_ID,Sample_Project\n" +
            "".join(f"{sid},{pid}\n" for sid, pid in selected.items()))
        book = load_workbook(run / "Sample-Submission-Form.xlsx")
        customer, lab = book["Sample-Submission-Form"], book["INFO (GCF-lab only)"]
        customer.delete_rows(16, customer.max_row)
        lab.delete_rows(2, lab.max_row)
        manifests = {}
        for row, (sid, pid) in enumerate(selected.items(), 16):
            for col, value in enumerate([sid, "007", "control", pid, 10], 1):
                customer.cell(row, col, value)
            lab.append([sid, 11 + number])
            (run / pid).mkdir(exist_ok=True)
            for mate in (1, 2):
                filename = f"{sid}_R{mate}.fastq.gz"
                data = gzip.compress(gzip.decompress((FIXTURE / RUN / PROJECT / filename).read_bytes()), mtime=number + 1)
                (run / pid / filename).write_bytes(data)
                manifests.setdefault(pid, []).append(f"{hashlib.md5(data).hexdigest()}  {pid}/{filename}\n")
        book.save(run / "Sample-Submission-Form.xlsx")
        for pid, lines in manifests.items():
            (run / f"md5sum_{pid}_fastq.txt").write_text("".join(lines))
        runs.append(run)
    work = tmp_path / "analysis"
    shutil.copytree(FIXTURE / "workflow", work / "src/gcf-workflows")
    return root, runs, work


def invoke(runs, work, projects, *options):
    return subprocess.run([str(Path(sys.executable).with_name("configmaker.py")),
                           *map(str, runs), "-p", *projects, *options],
                          cwd=work, capture_output=True, text=True, timeout=30)


def read_rows(path):
    with path.open() as handle:
        return list(csv.DictReader(handle))


def assert_provenance(root, work, assignments, override, batch):
    config = yaml.safe_load((work / "config.yaml").read_text())
    pep = yaml.safe_load((work / "pep/pep_config.yaml").read_text())
    expected_projects = list(dict.fromkeys(pid for run in assignments for pid in run.values()))
    assert config["project_id"] == pep["project_id"] == ([ANALYSIS] if override else expected_projects)
    if override:
        assert config["src_project_id"] == pep["src_project_id"] == expected_projects
    expected = {}
    for number, selected in enumerate(assignments):
        run = f"99010{number + 1}_SYNTHETIC_000{number + 1}_ABATCH{number + 1}"
        for sid, pid in selected.items():
            effective = f"{sid}_ABATCH{number + 1}" if batch else sid
            expected.setdefault(effective, []).append((run, pid, sid))
    assert set(config["samples"]) == set(expected)
    samples = {row["sample_name"]: row for row in read_rows(work / "pep/sample_table.csv")}
    subsamples = read_rows(work / "pep/subsample_table.csv") if "subsample_table" in pep else []
    for sid, contributions in expected.items():
        entry = config["samples"][sid]
        assert entry["Sample_ID"] == sid
        assert entry["Concentration"] == str(float(int(contributions[-1][0].split("_")[2]) + 10))
        assert entry["Project_ID"] == (ANALYSIS if override else ",".join(p for _, p, _ in contributions))
        if override:
            assert entry["Src_Project_ID"] == ",".join(p for _, p, _ in contributions)
        if batch:
            assert entry["Src_Sample_ID"] == contributions[0][2]
        rows = [dict(samples[sid], **row) for row in subsamples if row["sample_name"] == sid] or [samples[sid]]
        for mate in (1, 2):
            key = f"R{mate}"
            paths = [f"{run}/{pid}/{raw}_{key}.fastq.gz" for run, pid, raw in contributions]
            assert entry[key].split(",") == paths
            # Resolve the PEP independently, including the sample/subsample join.
            resolved = [pep["sample_modifiers"]["derive"]["sources"][row[key]].format(**row) for row in rows]
            assert resolved == paths
            digests = []
            for relative in paths:
                link = work / config["fastq_dir"] / relative
                assert link.is_file() and link.resolve() == (root / relative).resolve()
                digests.append(hashlib.md5((root / relative).read_bytes()).hexdigest())
            assert entry[key + "_md5sum"].split(",") == digests
            assert [row[key + "_md5sum"] for row in rows] == digests
    return config, pep


@pytest.mark.parametrize("override", [False, True])
@pytest.mark.parametrize("title", [None, "Combined synthetic experiment"])
def test_single_project_identity_and_text_title(tmp_path, override, title):
    assignments = [{"001": PROJECT}]
    root, runs, work = make_sources(tmp_path, assignments)
    before = contents(root)
    options = (["-P", ANALYSIS] if override else []) + (["--title", title] if title else [])
    result = invoke(runs, work, [PROJECT], *options)
    assert result.returncode == 0, result.stdout + result.stderr
    config, pep = assert_provenance(root, work, assignments, override, False)
    assert config["experiment_title"] == pep["experiment_title"] == (title or (ANALYSIS if override else PROJECT))
    assert contents(root) == before


@pytest.mark.parametrize("assignments", [
    [{"001": PROJECT, "001A": SECOND, "CTRL": PROJECT}],
    [{"001": PROJECT}, {"001A": SECOND}],
    [{"001": PROJECT, "CTRL": SECOND}, {"001": SECOND, "001A": PROJECT}],
    [{"001": PROJECT}, {"001": PROJECT}],
])
@pytest.mark.parametrize("override", [False, True])
@pytest.mark.parametrize("batch", [False, True])
def test_combined_sources_paths_checksums_and_metadata(tmp_path, assignments, override, batch):
    root, runs, work = make_sources(tmp_path, assignments)
    before = contents(root)
    projects = list(dict.fromkeys(pid for selected in assignments for pid in selected.values()))
    result = invoke(runs, work, projects, *(["-P", ANALYSIS] if override else []), *(["--keep-batch"] if batch else []))
    assert result.returncode == 0, result.stdout + result.stderr
    config, pep = assert_provenance(root, work, assignments, override, batch)
    assert config["experiment_title"] == pep["experiment_title"] == ", ".join([ANALYSIS] if override else projects)
    assert contents(root) == before


@pytest.mark.parametrize("unusable", ["absent", "empty", "no_selected_metadata"])
def test_explicit_requested_source_cannot_disappear(tmp_path, unusable):
    root, runs, work = make_sources(tmp_path, [{"001": PROJECT}])
    if unusable != "absent":
        (runs[0] / SECOND).mkdir()
    if unusable == "empty":
        sheet = runs[0] / "SampleSheet.csv"
        sheet.write_text(sheet.read_text() + f"001A,{SECOND}\n")
        book = load_workbook(runs[0] / "Sample-Submission-Form.xlsx")
        for col, value in enumerate(["001A", "008", "control", SECOND, 10], 1):
            book["Sample-Submission-Form"].cell(17, col, value)
        book["INFO (GCF-lab only)"].append(["001A", 11])
        book.save(runs[0] / "Sample-Submission-Form.xlsx")
    (work / "config.yaml").write_text("retained: keep\n")
    before, source_before = contents(work), contents(root)
    result = invoke(runs, work, [PROJECT, SECOND], "-P", ANALYSIS)
    assert result.returncode == 2
    assert SECOND in result.stderr
    assert "requested source" in result.stderr.lower()
    assert contents(work) == before
    assert contents(root) == source_before


def test_partial_discovery_keeps_observed_source_assignment(tmp_path):
    assignments = [{"001": PROJECT}, {"001": SECOND, "001A": SECOND}]
    root, runs, work = make_sources(tmp_path, assignments)
    for mate in (1, 2):
        (runs[1] / SECOND / f"001_R{mate}.fastq.gz").unlink()
    result = invoke(runs, work, [PROJECT, SECOND], "-P", ANALYSIS)
    assert result.returncode == 0, result.stdout + result.stderr
    config = yaml.safe_load((work / "config.yaml").read_text())
    assert config["samples"]["001"]["Src_Project_ID"] == PROJECT
    assert config["samples"]["001A"]["Src_Project_ID"] == SECOND
    assert "001" in result.stderr and "fastq" in result.stderr.lower()
    summary = json.loads((work / "configmaker.analysis-summary.json").read_text())
    assert summary["sample_count"] == summary["planned_sample_count"] == 2


@pytest.mark.parametrize("batch", [False, True])
@pytest.mark.parametrize("indexed", [False, True])
def test_discovered_lane_paths_are_not_reconstructed_from_analysis_names(tmp_path, batch, indexed):
    root, runs, work = make_sources(tmp_path, [{"001": PROJECT}])
    project = runs[0] / PROJECT
    expected = {read: [] for read in ("R1", "R2", *(["I1"] if indexed else []))}
    manifest = []
    for read in expected:
        data = (project / "001_R1.fastq.gz").read_bytes()
        for lane in ("L001", "L002"):
            path = project / f"001_S1_{lane}_{read}_001.fastq.gz"
            path.write_bytes(data)
            expected[read].append(str(path.relative_to(root)))
            manifest.append(f"{hashlib.md5(data).hexdigest()}  {PROJECT}/{path.name}\n")
    for mate in (1, 2):
        (project / f"001_R{mate}.fastq.gz").unlink()
    (runs[0] / f"md5sum_{PROJECT}_fastq.txt").write_text("".join(manifest))
    before = contents(root)
    result = invoke(runs, work, [PROJECT], "-P", ANALYSIS, *(["--keep-batch"] if batch else []))
    assert result.returncode == 0, result.stdout + result.stderr
    config = yaml.safe_load((work / "config.yaml").read_text())
    pep = yaml.safe_load((work / "pep/pep_config.yaml").read_text())
    sample = read_rows(work / "pep/sample_table.csv")[0]
    rows = [dict(sample, **row) for row in read_rows(work / "pep/subsample_table.csv")]
    entry = next(iter(config["samples"].values()))
    for read, paths in expected.items():
        assert entry[read].split(",") == paths
        assert [pep["sample_modifiers"]["derive"]["sources"][row[read]].format(**row) for row in rows] == paths
        assert entry[read + "_md5sum"].split(",") == [row[read + "_md5sum"] for row in rows]
        for path in paths:
            assert (work / config["fastq_dir"] / path).resolve() == root / path
    assert contents(root) == before
