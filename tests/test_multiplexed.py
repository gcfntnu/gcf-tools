import json

import pytest
from openpyxl import load_workbook

from configmaker.validation import validate_inputs
from multiplexed_fixture import CUSTOMER, LAB, DEMUX, KIT, make_multiplexed


@pytest.mark.parametrize("worksheet,row,header_row", [(CUSTOMER, 16, 15), (LAB, 2, 1)])
@pytest.mark.parametrize("header,field", [
    ("Sample Group [dtype=string]", "Sample_Group"),
    ("Sample biosource", "Sample_Biosource"),
    ("External ID (optional reference sample ID)", "External_ID"),
    ("Sample_Group", "Sample_Group"), ("Sample_Biosource", "Sample_Biosource"), ("External_ID", "External_ID"),
])
@pytest.mark.parametrize("value", ["biological", 0, False])
def test_reserved_fields_rejected_before_conversion(tmp_path, worksheet, row, header_row, header, field, value):
    run, form, _ = make_multiplexed(tmp_path)
    book = load_workbook(form)
    tab = book[worksheet]
    # Replace template biological column in customer to avoid header collisions.
    col = {"Sample_Group": 3, "External_ID": 2}.get(field, 12) if worksheet == CUSTOMER else 5
    tab.cell(header_row, col, header)
    tab.cell(row, col, value)
    book.save(form)
    before = form.read_bytes()
    result = validate_inputs(run / "SampleSheet.csv", form)
    errors = [f for f in result.to_dict()["errors"] if f["code"] == "workbook.pooled_biological_metadata"]
    assert len(errors) == 1, result.render_text()
    error = errors[0]
    assert (error["file"], error["worksheet"], error["row"], error["column"], error["sample_id"]) == (
        str(form), worksheet, row, field, "001")
    assert error["details"]["value"] == value
    assert "Cell Multiplexing" in error["message"]
    assert not result.ok and form.read_bytes() == before
    json.dumps(result.to_dict())


@pytest.mark.parametrize("value", [None, "", " ", "NA"])
def test_empty_reserved_cells_pass(tmp_path, value):
    run, form, _ = make_multiplexed(tmp_path)
    book = load_workbook(form)
    book[CUSTOMER]["B16"] = value
    book[CUSTOMER]["C16"] = value
    book[LAB].cell(1, 5, "Sample_Biosource [dtype=string]")
    book[LAB].cell(2, 5, value)
    book.save(form)
    result = validate_inputs(run / "SampleSheet.csv", form)
    assert result.ok, result.render_text()
    assert "External_ID" not in result.metadata
    assert result.forms[0]["demux"]["data"].Sample_ID.tolist() == ["001", "1"]


def test_invalid_values_cannot_be_hidden_by_join_or_later_form(tmp_path):
    run, bad, _ = make_multiplexed(tmp_path)
    good = bad.with_name("good.xlsx")
    good.write_bytes(bad.read_bytes())
    book = load_workbook(bad)
    book[CUSTOMER]["C16"] = "invalid"
    book[LAB].append(["UNJOINED", 1, None, None, "hidden"])
    book[LAB].cell(1, 5, "External_ID")
    book.save(bad)
    result = validate_inputs(run / "SampleSheet.csv", [bad, good])
    errors = [f for f in result.to_dict()["errors"] if f["code"] == "workbook.pooled_biological_metadata"]
    assert {(f["sample_id"], f["column"]) for f in errors} == {("001", "Sample_Group"), ("UNJOINED", "External_ID")}
    assert {f["file"] for f in errors} == {str(bad)}


@pytest.mark.parametrize("kit,override,invalid", [(KIT, None, True), ("Synthetic", KIT, True),
                                                  (KIT, "Synthetic", True), ("10x Test", None, False)])
def test_reserved_rule_uses_existing_parse_selection(tmp_path, kit, override, invalid):
    run, form, _ = make_multiplexed(tmp_path, sheet_kit=kit)
    book = load_workbook(form)
    book[CUSTOMER]["B16"] = "007"
    book[CUSTOMER]["C16"] = "control"
    book.save(form)
    result = validate_inputs(run / "SampleSheet.csv", form, libkit=override)
    assert result.ok is not invalid, result.render_text()
    if not invalid:
        assert result.metadata.loc["001", "External_ID"] == "007"
        assert result.metadata.loc["001", "Sample_Group"] == "control"
        assert "Volume" not in result.metadata and "KIT" not in result.metadata
        assert result.forms[0]["demux"] is None


@pytest.mark.parametrize("worksheet,row", [(CUSTOMER, 15), (LAB, 1), (DEMUX, 1)])
def test_custom_mapped_header_collision_still_fails(tmp_path, worksheet, row):
    run, form, _ = make_multiplexed(tmp_path)
    book = load_workbook(form)
    tab = book[worksheet]
    column = tab.max_column + 1
    tab.cell(row, column, "Novel Field")
    # Descriptor stripping must still detect duplicate mapped headers.
    tab.cell(row, column + 1, "Novel Field [dtype=string]")
    tab.cell(row + 1, column, 0)
    tab.cell(row + 1, column + 1, 1)
    book.save(form)
    result = validate_inputs(run / "SampleSheet.csv", form)
    assert any(f["code"] == "workbook.duplicate_columns" and f["worksheet"] == worksheet
               for f in result.to_dict()["errors"])
