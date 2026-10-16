"""Offline installed initialization for global and descriptive organism values."""
import csv
from pathlib import Path
import shutil
import subprocess
import sys

from openpyxl import load_workbook
import pytest
import yaml

from test_paired_cli import FIXTURE, PROJECT, RUN, SAMPLES, contents


def prepare(tmp_path, global_value='Homo sapiens', annotations=None):
    source=tmp_path/'source'
    shutil.copytree(FIXTURE, source)
    sheet=source/RUN/'SampleSheet.csv'
    text=sheet.read_text().replace('Organism,Homo sapiens\n', '' if global_value is None else 'Organism,'+global_value+'\n')
    sheet.write_text(text)
    if annotations is not None:
        path=source/RUN/'Sample-Submission-Form.xlsx'
        book=load_workbook(path);customer=book['Sample-Submission-Form'];customer.cell(15,6,'Organism')
        for row,value in enumerate(annotations,16):customer.cell(row,6,value)
        book.save(path)
    work=tmp_path/'analysis'
    shutil.copytree(source/'workflow',work/'src/gcf-workflows')
    return source,work


def invoke(source,work,*options):
    return subprocess.run([str(Path(sys.executable).with_name('configmaker.py')),str(source/RUN),'-p',PROJECT,*options],
                          cwd=work,capture_output=True,text=True,timeout=30)


@pytest.mark.parametrize('sheet,cli,expected', [('Human',None,'homo_sapiens'),('Mouse',' Human ','homo_sapiens'),
    ('hOmO SaPiEnS',None,'homo_sapiens'),('Human','Mus musculus','mus_musculus'),
    ('rangifer_tarandus',None,'rangifer_tarandus'),('Human',' Rangifer tarandus ','rangifer_tarandus'),
    ('homo_sapiens__mus_musculus',None,'homo_sapiens__mus_musculus'),
    ('Human',' Human __ Mouse ','homo_sapiens__mus_musculus'),
    ('N/A',None,None),('Human','N/A','N/A'),(None,None,None)])
def test_effective_global_value_normalizes_without_cache(tmp_path,sheet,cli,expected):
    source,work=prepare(tmp_path,sheet)
    before=contents(source)
    result=invoke(source,work,*(['--organism',cli] if cli is not None else []))
    assert result.returncode==0,result.stdout+result.stderr
    config=yaml.safe_load((work/'config.yaml').read_text());pep=yaml.safe_load((work/'pep/pep_config.yaml').read_text())
    assert config.get('organism')==pep.get('organism')==expected
    assert contents(source)==before


@pytest.mark.parametrize('value',['Homo sapens','unlisted synthetic taxon','reference','Human__Homo sapens','Human__reference'])
def test_unresolved_required_global_preserves_every_output(tmp_path,value):
    source,work=prepare(tmp_path)
    (work/'config.yaml').write_text('retained: keep\n')
    (work/'data').mkdir();(work/'data/sentinel').write_text('keep\n')
    before,inputs=contents(work),contents(source)
    result=invoke(source,work,'--organism',value)
    assert result.returncode==2
    assert 'organism' in result.stderr.lower() and value in result.stderr
    assert 'canonical' in result.stderr.lower()
    assert contents(work)==before
    assert contents(source)==inputs


@pytest.mark.parametrize('global_value',[None,'N/A','Homo sapiens'])
def test_mixed_species_and_unresolved_annotations_keep_their_roles(tmp_path,global_value):
    unknown='  Unlisted annotation (retain verbatim)!  '
    source,work=prepare(tmp_path,global_value,['Human','Mouse',unknown])
    result=invoke(source,work)
    assert result.returncode==0,result.stdout+result.stderr
    config=yaml.safe_load((work/'config.yaml').read_text())
    assert config.get('organism')==({'Homo sapiens':'homo_sapiens','N/A':None}.get(global_value,global_value))
    assert [config['samples'][sid]['Organism'] for sid in SAMPLES]==['homo_sapiens','mus_musculus',unknown]
    with (work/'pep/sample_table.csv').open() as handle:rows=list(csv.DictReader(handle))
    assert [row['sample_name'] for row in rows]==list(SAMPLES)
    assert [row['Organism'] for row in rows]==['homo_sapiens','mus_musculus',unknown]
    assert 'CTRL' in result.stderr and 'Organism' in result.stderr and 'unresolved' in result.stderr.lower()


def test_uniform_unresolved_annotation_is_not_erased_by_global_setting(tmp_path):
    unknown='  Unlisted annotation!  '
    source,work=prepare(tmp_path,'Human',[unknown]*3)
    result=invoke(source,work)
    assert result.returncode==0,result.stdout+result.stderr
    config=yaml.safe_load((work/'config.yaml').read_text())
    assert config['organism']=='homo_sapiens'
    assert [s['Organism'] for s in config['samples'].values()]==[unknown]*3


@pytest.mark.parametrize('global_value',[None,'Human'])
def test_uniform_mixed_annotation_is_not_collapsed_to_a_global_species(tmp_path,global_value):
    source,work=prepare(tmp_path,global_value,['Human__Mouse']*3)
    result=invoke(source,work)
    assert result.returncode==0,result.stdout+result.stderr
    config=yaml.safe_load((work/'config.yaml').read_text())
    assert config.get('organism')==('homo_sapiens' if global_value else None)
    assert [s['Organism'] for s in config['samples'].values()]==['homo_sapiens__mus_musculus']*3


@pytest.mark.parametrize('workflow',['metagenome','rnaseq'])
@pytest.mark.parametrize('explicit_snapshot',[False,True])
def test_established_community_label_is_not_a_species_reference(tmp_path,workflow,explicit_snapshot):
    source,work=prepare(tmp_path,' MeTaGeNoMe ')
    path=work/'src/gcf-workflows/libprep.config'
    selected=path.read_text().replace('workflow: rnaseq','workflow: '+workflow)
    if explicit_snapshot:
        path=tmp_path/'selected-libprep.config'
    path.write_text(selected)
    before=contents(work),contents(source)
    result=invoke(source,work,*(['--libprep-config',str(path)] if explicit_snapshot else []))
    assert path.read_text()==selected
    if workflow=='rnaseq':
        assert result.returncode==2 and 'Project organism' in result.stderr
        assert (contents(work),contents(source))==before
    else:
        assert result.returncode==0,result.stdout+result.stderr
        config=yaml.safe_load((work/'config.yaml').read_text())
        pep=yaml.safe_load((work/'pep/pep_config.yaml').read_text())
        assert config['organism']==pep['organism']=='metagenome'
        assert config['workflow']=='metagenome'
        assert config['libprep_selection']['sha256']
