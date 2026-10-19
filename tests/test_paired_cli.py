"""Installed default-PEP contract, with an independent paired-record oracle."""
import csv
import gzip
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys

import pytest
import yaml
from openpyxl import load_workbook

FIXTURE = Path(__file__).parent / "fixtures/paired"
RUN = "990101_SYNTHETIC_0001_ASYNTHETIC"
PROJECT = "GCF-2099-001"
SAMPLES = ("001", "001A", "CTRL")
EXPECTED = {
    "001": {"External_ID": "007", "Sample_Group": "control", "Concentration": "11.0"},
    "001A": {"External_ID": "008", "Sample_Group": "treated", "Concentration": "22.0"},
    "CTRL": {"External_ID": "009", "Sample_Group": "control", "Concentration": "33.0"},
}


def contents(root):
    return {str(p.relative_to(root)): p.read_bytes() for p in root.rglob("*") if p.is_file()}


def assert_pairs(first, second, sample):
    """No producer/configmaker imports; validate records, order and identity."""
    sequences = []
    for mate, path, length in [(1, first, 75), (2, second, 150)]:
        lines = gzip.decompress(path.read_bytes()).decode("ascii").splitlines()
        assert len(lines) == 24 * 4
        observed = []
        for index in range(24):
            header, sequence, plus, quality = lines[index * 4:index * 4 + 4]
            assert header == f"@synthetic:{sample}:{index + 1:04d}/{mate}"
            assert plus == "+"
            assert len(sequence) == len(quality) == length
            assert set(sequence) <= set("ACGT")
            assert all(33 <= ord(char) <= 126 for char in quality)
            observed.append(sequence)
        assert len(set(observed)) == 24
        sequences.append(observed)
    assert sequences[0] != [sequence[:75] for sequence in sequences[1]]


def assert_inputs(root):
    provenance = json.loads((root / "provenance.json").read_text())
    actual = {name: hashlib.sha256(data).hexdigest() for name, data in contents(root).items()
              if name != "provenance.json"}
    assert actual == provenance["sha256"]
    run = root / RUN
    manifest = dict(line.split() for line in (run / f"md5sum_{PROJECT}_fastq.txt").read_text().splitlines())
    assert len(manifest) == 6
    for sample in SAMPLES:
        first, second = [run / PROJECT / f"{sample}_R{mate}.fastq.gz" for mate in (1, 2)]
        assert_pairs(first, second, sample)
        for path in (first, second):
            assert manifest[hashlib.md5(path.read_bytes()).hexdigest()] == f"{PROJECT}/{path.name}"
    book = load_workbook(run / "Sample-Submission-Form.xlsx")
    customer = book["Sample-Submission-Form"]
    assert [customer.cell(row, 1).value for row in range(16, 19)] == list(SAMPLES)
    assert [customer.cell(row, 5).value for row in range(16, 19)] == [10, 20, 30]
    assert [row for row in book["INFO (GCF-lab only)"].values] == [
        ("Sample_ID", "Concentration"), ("001", 11), ("001A", 22), ("CTRL", 33)]


@pytest.fixture
def example(tmp_path):
    source = tmp_path / "input"
    shutil.copytree(FIXTURE, source)
    work = tmp_path / "analysis"
    shutil.copytree(source / "workflow", work / "src/gcf-workflows")
    return source, work


def cli(source, work):
    command = [str(Path(sys.executable).parent / "configmaker.py"), str(source / RUN),
               "--project-id", PROJECT]
    result = subprocess.run(command, cwd=work, capture_output=True, text=True, timeout=30)
    # Retain diagnostics outside the output tree, including failed examples.
    (work.parent / "cli-transcript.txt").write_text(result.stdout + result.stderr)
    return result


def assert_outputs(source, work, samples=SAMPLES):
    config = yaml.safe_load((work / "config.yaml").read_text())
    assert config["project_id"] == [PROJECT]
    assert config["experiment_title"] == PROJECT
    assert config["organism"] == "homo_sapiens"
    assert config["read_geometry"] == [75, 150]
    assert all(type(length) is int for length in config["read_geometry"])
    assert config["workflow"] == "rnaseq"
    assert config["library_layout"] == "paired"
    assert config["libprepkit"] == "Synthetic"
    assert config["filter"]["subsample_fastq"] == "skip"
    assert config["filter"]["trim"]["fastp"]["params"] == "-q 17"
    assert config["quant"]["batch"] == {"method": "skip"}
    assert config["multiple_flowcells"] is False
    assert config["fastq_dir"] == "data/raw/fastq"
    kit = (source / "workflow/libprep.config").read_bytes()
    assert (work / "src/gcf-workflows/libprep.config").read_bytes() == kit
    selection = config["libprep_selection"]
    assert selection["sha256"] == hashlib.sha256(kit).hexdigest()
    assert selection["entry"] == "Synthetic PE"
    assert selection["workflow"] == "rnaseq"
    assert selection["read_geometry"] == [75, 150]
    assert set(config["samples"]) == set(samples)
    pep = yaml.safe_load((work / "pep/pep_config.yaml").read_text())
    assert pep["pep_version"] == "2.0.0"
    assert pep["project_id"] == [PROJECT]
    assert pep["read_geometry"] == [75, 150]
    assert pep["workflow"] == "rnaseq"
    assert pep["sample_table"] == "sample_table.csv"
    assert "subsample_table" not in pep
    assert not (work / "pep/subsample_table.csv").exists()
    with (work / "pep/sample_table.csv").open() as handle:
        rows = {row["sample_name"]: row for row in csv.DictReader(handle)}
    assert set(rows) == set(samples)
    for sample in samples:
        entry, row = config["samples"][sample], rows[sample]
        assert entry["Sample_ID"] == sample
        assert entry["Project_ID"] == row["Project_ID"] == PROJECT
        assert entry["Flowcell_Name"] == row["Flowcell_Name"] == RUN
        assert entry["Flowcell_ID"] == row["Flowcell_ID"] == "ASYNTHETIC"
        for key, value in EXPECTED[sample].items():
            assert entry[key] == row[key] == value
            assert type(entry[key]) is str
        for mate in (1, 2):
            key = f"R{mate}"
            relative = f"{RUN}/{PROJECT}/{sample}_{key}.fastq.gz"
            assert entry[key] == relative
            assert row[key] in pep["sample_modifiers"]["derive"]["sources"]
            # Evaluate the published PEP source template independently of exporter.
            assert pep["sample_modifiers"]["derive"]["sources"][row[key]].format(**row) == relative
            link = work / config["fastq_dir"] / relative
            original = source / relative
            assert link.is_symlink() and link.is_file()
            assert link.resolve() == original.resolve()
            digest = hashlib.md5(original.read_bytes()).hexdigest()
            assert entry[key + "_md5sum"] == row[key + "_md5sum"] == digest
        assert_pairs(work / config["fastq_dir"] / entry["R1"],
                     work / config["fastq_dir"] / entry["R2"], sample)
    assert len(list((work / "data/raw/fastq").rglob("*.fastq.gz"))) == len(samples) * 2
    snakefile = (work / "Snakefile").read_text()
    for expected in ('config.yaml', 'pep/pep_config.yaml', 'src/gcf-workflows/rnaseq/rnaseq.smk'):
        assert expected in snakefile
    assert (work / "src/gcf-workflows/rnaseq/rnaseq.smk").is_file()
    summary = json.loads((work / "configmaker.analysis-summary.json").read_text())
    assert summary == {
        "schema_version": 1, "kind": "fastq_discovery", "sample_count": len(samples),
        "planned_sample_count": 3, "missing_sample_ids": sorted(set(SAMPLES) - set(samples)),
        "samples": [{"sample_id": sample, "flowcells": [RUN]} for sample in samples],
        "project_ids": [PROJECT], "read_geometry": [75, 150], "run_read_geometry": [75, 150],
        "metadata_inputs": [
            {"kind": kind, "path": str(source / RUN / filename),
             "sha256": hashlib.sha256((source / RUN / filename).read_bytes()).hexdigest(),
             "size_bytes": (source / RUN / filename).stat().st_size}
            for kind, filename in (("samplesheet", "SampleSheet.csv"),
                                   ("submission_form", "Sample-Submission-Form.xlsx"))
        ],
        "samplesheet_options": {"Libprep": "Synthetic", "Organism": "Homo sapiens"},
        "libprep_selection": {
            "source": str(work / "src/gcf-workflows/libprep.config"),
            "sha256": hashlib.sha256(kit).hexdigest(), "kit": "Synthetic", "entry": "Synthetic PE",
            "read_geometry": [75, 150], "workflow": "rnaseq",
        },
    }
    assert f"FASTQ discovery: {len(samples)} of 3 planned" in (work / ".configmaker.log").read_text()


def test_paired_example_installed_cli_default_pep(example):
    source, work = example
    assert_inputs(source)
    before = contents(source)
    result = cli(source, work)
    assert result.returncode == 0, result.stdout + result.stderr
    assert_outputs(source, work)
    assert contents(source) == before


@pytest.mark.parametrize("qualified", [False, True])
def test_canonical_legacy_pair_ignores_unselected_metadata_backups(example, qualified):
    source, work = example
    run = source / RUN
    sheet = run / ((PROJECT + "_" if qualified else "") + "SampleSheet_alternate.csv")
    sheet.write_text("An unrelated historical backup,not selected\n")
    (run / "Sample-Submission-Form_backup.xlsx").write_bytes(b"Unrelated backup, not selected")
    before = contents(source)
    result = cli(source, work)
    assert result.returncode == 0, result.stdout + result.stderr
    assert_outputs(source, work)
    assert contents(source) == before


def test_paired_fixture_regeneration(tmp_path):
    from paired_fixture import create_fixture
    regenerated = create_fixture(tmp_path / "regenerated")
    assert contents(regenerated) == contents(FIXTURE)
    assert_inputs(regenerated)


@pytest.mark.parametrize("existing", [False, True])
def test_paired_invalid_metadata_preserves_output(example, existing):
    source, work = example
    form = source / RUN / "Sample-Submission-Form.xlsx"
    book = load_workbook(form)
    book["Sample-Submission-Form"]["A16"] = "WRONG"
    book.save(form)
    if existing:
        (work / "config.yaml").write_text("existing: keep\n")
        (work / "data").mkdir()
        (work / "data/sentinel").write_text("keep\n")
    before, inputs = contents(work), contents(source)
    result = cli(source, work)
    assert result.returncode == 2
    assert "sample.merge_loss" in result.stderr
    assert contents(work) == before
    assert contents(source) == inputs


@pytest.mark.parametrize("remaining", [(), ("001A", "CTRL")])
def test_paired_empty_and_partial_discovery(example, remaining):
    source, work = example
    for sample in set(SAMPLES) - set(remaining):
        for mate in (1, 2):
            (source / RUN / PROJECT / f"{sample}_R{mate}.fastq.gz").unlink()
    before, inputs = contents(work), contents(source)
    result = cli(source, work)
    if remaining:
        assert result.returncode == 0, result.stdout + result.stderr
        assert_outputs(source, work, remaining)
    else:
        assert result.returncode == 2
        assert "FASTQ discovery failed" in result.stderr
        assert contents(work) == before
    assert contents(source) == inputs


@pytest.mark.parametrize("damage", ["mate_id", "count", "quality"])
def test_pair_oracle_rejects_corruption(example, damage):
    source, _ = example
    first = source / RUN / PROJECT / "001_R1.fastq.gz"
    second = source / RUN / PROJECT / "001_R2.fastq.gz"
    lines = gzip.decompress(second.read_bytes()).decode().splitlines()
    if damage == "mate_id":
        lines[0] = "@synthetic:001A:0001/2"
    elif damage == "count":
        lines = lines[:-4]
    else:
        lines[3] = lines[3][:-1]
    second.write_bytes(gzip.compress(("\n".join(lines) + "\n").encode(), mtime=0))
    with pytest.raises(AssertionError):
        assert_pairs(first, second, "001")
