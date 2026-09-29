"""Validate real SampleSheet/workbook combinations before config construction."""

import csv
import hashlib
import io
import json

import pytest
from openpyxl import Workbook

from configmaker.validation import validate_inputs


CUSTOMER_HEADERS = [
    "Unique Sample ID",
    "External ID (optional reference sample ID)",
    "Sample Group",
    "Project ID",
]
LAB_HEADERS = ["Sample_ID", "Concentration"]
PROJECT = "GCF-2026-001"


def sample_sheet(path, rows, headers=None):
    """Create an IEM SampleSheet, keeping exact identifiers including whitespace."""
    path.parent.mkdir(parents=True, exist_ok=True)
    stream = io.StringIO()
    stream.write("[Header]\nIEMFileVersion,4\n[Data]\n")
    writer = csv.writer(stream)
    writer.writerow(headers or ["Sample_ID", "Sample_Project"])
    writer.writerows(rows)
    path.write_text(stream.getvalue())
    return path


def submission_form(
    path,
    customer_rows,
    lab_rows=(),
    customer_headers=None,
    lab_headers=None,
    include_customer=True,
    include_lab=True,
    formatted_empty_rows=False,
):
    """Build the production two-sheet layout without relying on checked-in data."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    if include_customer:
        customer = workbook.create_sheet("Sample-Submission-Form")
        customer.cell(1, 1, "GCF sample submission form")
        for index, header in enumerate(customer_headers or CUSTOMER_HEADERS, 1):
            customer.cell(15, index, header)
        for row_index, row in enumerate(customer_rows, 16):
            for column_index, value in enumerate(row, 1):
                customer.cell(row_index, column_index, value)
        if formatted_empty_rows:
            customer.cell(100, 1).number_format = "@"
            customer.cell(100, 2, "")
    if include_lab:
        lab = workbook.create_sheet("INFO (GCF-lab only)")
        lab.append(lab_headers or LAB_HEADERS)
        for row in lab_rows:
            lab.append(row)
        if formatted_empty_rows:
            lab.cell(100, 1).number_format = "@"
    if not workbook.worksheets:
        workbook.create_sheet("Unrelated sheet")
    workbook.save(path)
    return path


def customer_row(sample_id, external_id=None, group="A", project=PROJECT):
    return [sample_id, external_id or sample_id, group, project]


def report_for(sheets, forms, **kwargs):
    result = validate_inputs([str(path) for path in sheets], [str(path) for path in forms], **kwargs)
    report = result.to_dict()
    assert result.ok is (len(report["errors"]) == 0)
    assert isinstance(result.render_text(), str)
    # JSON output must work without a custom encoder, including Excel values.
    json.dumps(report)
    for severity, plural in [("error", "errors"), ("warning", "warnings"), ("info", "info")]:
        for finding in report[plural]:
            assert finding["severity"] == severity
            assert isinstance(finding["code"], str) and finding["code"]
            assert isinstance(finding["message"], str) and finding["message"]
            if "file" in finding:
                assert isinstance(finding["file"], str)
            if "row" in finding:
                assert isinstance(finding["row"], int)
            if "details" in finding:
                assert isinstance(finding["details"], dict)
    return result, report


def diagnostics(report, category="errors"):
    return json.dumps(report[category], ensure_ascii=False).lower()


def assert_error_mentions(report, *terms):
    assert report["errors"], report
    combined = diagnostics(report)
    for term in terms:
        assert str(term).lower() in combined, report["errors"]


def test_valid_subset_reports_planned_samples_and_unused_submission_rows(tmp_path):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1", PROJECT]])
    form = submission_form(
        tmp_path / "submission.xlsx", [customer_row("S1"), customer_row("S2")]
    )
    result, report = report_for([sheet], [form])
    assert result.ok
    assert report["summary"]["planned_sample_count"] == 1
    assert "S1" in json.dumps(report["summary"]["planned_samples"])
    assert PROJECT in json.dumps(report["summary"]["projects"])
    assert "A" in json.dumps(report["summary"]["sample_groups"])
    assert "s2" in diagnostics(report, "info")
    assert not report["warnings"]


def test_report_fingerprints_exact_files_without_changing_them(tmp_path):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1", PROJECT]])
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("S1")])
    originals = {str(path): path.read_bytes() for path in [sheet, form]}
    _, report = report_for([sheet], [form])
    by_path = {item["path"]: item for item in report["inputs"]}
    assert set(by_path) == set(originals)
    for path in [sheet, form]:
        item = by_path[str(path)]
        assert item["sha256"] == hashlib.sha256(originals[str(path)]).hexdigest()
        assert item["size_bytes"] == len(originals[str(path)])
        assert item["kind"]
        assert path.read_bytes() == originals[str(path)]


def test_missing_submission_sample_is_blocking_and_identifies_source(tmp_path):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1", PROJECT], ["S2", PROJECT]])
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("S1")])
    result, report = report_for([sheet], [form])
    assert not result.ok
    assert_error_mentions(report, "S2")
    assert "S2" in result.render_text()


def test_customer_lab_inner_join_loss_names_sample_and_workbook(tmp_path):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1", PROJECT], ["S2", PROJECT]])
    form = submission_form(
        tmp_path / "submission.xlsx",
        [customer_row("S1"), customer_row("S2")],
        [["S1", 4.5]],
    )
    _, report = report_for([sheet], [form])
    assert_error_mentions(report, "S2", "lab", str(form))


def test_lab_only_requested_sample_does_not_pass_customer_lab_join(tmp_path):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S2", PROJECT]])
    form = submission_form(
        tmp_path / "submission.xlsx", [customer_row("S1")], [["S1", 4.5], ["S2", 3]]
    )
    _, report = report_for([sheet], [form])
    assert_error_mentions(report, "S2", "customer")


def test_empty_lab_sheet_is_supported_and_blank_template_rows_are_ignored(tmp_path):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1", PROJECT]])
    form = submission_form(
        tmp_path / "submission.xlsx", [customer_row("S1")], formatted_empty_rows=True
    )
    result, report = report_for([sheet], [form])
    assert result.ok, report
    assert report["summary"]["planned_sample_count"] == 1


@pytest.mark.parametrize("sheet_name", ["customer", "lab"])
def test_populated_submission_row_without_id_is_not_silently_discarded(tmp_path, sheet_name):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1", PROJECT]])
    customer_rows = [customer_row("S1")]
    lab_rows = [["S1", 4.5]]
    if sheet_name == "customer":
        customer_rows.append([None, "external-2", "A", PROJECT])
    else:
        lab_rows.append([None, 3.1])
    form = submission_form(tmp_path / "submission.xlsx", customer_rows, lab_rows)
    _, report = report_for([sheet], [form])
    assert_error_mentions(report, "id", str(form))


@pytest.mark.parametrize("sheet_name", ["customer", "lab"])
def test_missing_required_workbook_sheet_is_actionable(tmp_path, sheet_name):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1", PROJECT]])
    form = submission_form(
        tmp_path / "submission.xlsx",
        [customer_row("S1")],
        include_customer=sheet_name != "customer",
        include_lab=sheet_name != "lab",
    )
    _, report = report_for([sheet], [form])
    expected_name = "Sample-Submission-Form" if sheet_name == "customer" else "INFO (GCF-lab only)"
    assert_error_mentions(report, expected_name, str(form))


@pytest.mark.parametrize("headers", [["Sample_ID"], ["Sample_Project"]])
def test_samplesheet_required_header_missing(tmp_path, headers):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1"]], headers=headers)
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("S1")])
    _, report = report_for([sheet], [form])
    missing_header = "Sample_Project" if headers == ["Sample_ID"] else "Sample_ID"
    assert_error_mentions(report, missing_header, str(sheet))


def test_missing_data_section_is_a_finding_instead_of_traceback(tmp_path):
    sheet = tmp_path / "SampleSheet.csv"
    sheet.write_text("[Header]\nIEMFileVersion,4\n")
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("S1")])
    _, report = report_for([sheet], [form])
    assert_error_mentions(report, "data", str(sheet))


@pytest.mark.parametrize("row,missing", [(["", PROJECT], "Sample_ID"), (["S1", ""], "Sample_Project")])
def test_samplesheet_populated_row_requires_identifiers(tmp_path, row, missing):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [row])
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("S1")])
    _, report = report_for([sheet], [form])
    assert_error_mentions(report, missing, str(sheet))


@pytest.mark.parametrize("source", ["samplesheet", "customer", "lab"])
def test_duplicate_headers_are_rejected_before_reader_mangles_names(tmp_path, source):
    sheet_headers = ["Sample_ID", "Sample_Project"]
    customer_headers = list(CUSTOMER_HEADERS)
    lab_headers = list(LAB_HEADERS)
    sheet_rows = [["S1", PROJECT]]
    customer_rows = [customer_row("S1")]
    lab_rows = [["S1", 4.5]]
    if source == "samplesheet":
        sheet_headers.append("Sample_ID")
        sheet_rows[0].append("S2")
    elif source == "customer":
        customer_headers.append("Unique Sample ID")
        customer_rows[0].append("S2")
    else:
        lab_headers.append("Sample_ID")
        lab_rows[0].append("S2")
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", sheet_rows, headers=sheet_headers)
    form = submission_form(
        tmp_path / "submission.xlsx", customer_rows, lab_rows,
        customer_headers=customer_headers, lab_headers=lab_headers,
    )
    _, report = report_for([sheet], [form])
    assert_error_mentions(report, "duplicate")


def test_different_headers_mapping_to_same_column_are_rejected(tmp_path):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1", PROJECT]])
    form = submission_form(
        tmp_path / "submission.xlsx",
        [customer_row("S1") + ["S2"]],
        customer_headers=CUSTOMER_HEADERS + ["Unique alternate ID"],
    )
    _, report = report_for([sheet], [form])
    assert_error_mentions(report, "Sample_ID")


@pytest.mark.parametrize("source", ["customer", "lab"])
def test_duplicate_submission_sample_ids_are_ambiguous(tmp_path, source):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1", PROJECT]])
    customer_rows = [customer_row("S1")]
    lab_rows = [["S1", 4.5]]
    if source == "customer":
        customer_rows.append(customer_row("S1", external_id="other"))
    else:
        lab_rows.append(["S1", 5.5])
    form = submission_form(tmp_path / "submission.xlsx", customer_rows, lab_rows)
    _, report = report_for([sheet], [form])
    assert_error_mentions(report, "duplicate", "S1")


def test_same_sample_in_distinct_lanes_is_one_planned_sample(tmp_path):
    sheet = sample_sheet(
        tmp_path / "SampleSheet.csv", [["S1", PROJECT, 1], ["S1", PROJECT, 2]],
        headers=["Sample_ID", "Sample_Project", "Lane"],
    )
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("S1")])
    result, report = report_for([sheet], [form])
    assert result.ok, report
    assert report["summary"]["planned_sample_count"] == 1


@pytest.mark.parametrize("lane", [None, 1])
def test_repeated_sample_in_same_lane_or_without_lanes_is_ambiguous(tmp_path, lane):
    headers = ["Sample_ID", "Sample_Project"] + (["Lane"] if lane else [])
    row = ["S1", PROJECT] + ([lane] if lane else [])
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [row, row], headers=headers)
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("S1")])
    _, report = report_for([sheet], [form])
    assert_error_mentions(report, "S1")


def test_lane_repetition_cannot_hide_conflicting_sample_metadata(tmp_path):
    sheet = sample_sheet(
        tmp_path / "SampleSheet.csv", [["S1", PROJECT, 1, "AAAA"], ["S1", PROJECT, 2, "CCCC"]],
        headers=["Sample_ID", "Sample_Project", "Lane", "Sample_Name"],
    )
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("S1")])
    _, report = report_for([sheet], [form])
    assert_error_mentions(report, "S1")


def test_leading_zero_identifiers_remain_distinct(tmp_path):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["001", PROJECT], ["1", PROJECT]])
    form = submission_form(
        tmp_path / "submission.xlsx", [customer_row("001"), customer_row("1")],
        [["001", 4.5], ["1", 5.5]],
    )
    result, report = report_for([sheet], [form])
    assert result.ok, report
    assert report["summary"]["planned_sample_count"] == 2
    assert "001" in json.dumps(report["summary"]["planned_samples"])


def test_whitespace_mismatch_is_explained_without_repairing_the_id(tmp_path):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1", PROJECT]])
    form = submission_form(tmp_path / "submission.xlsx", [customer_row(" S1 ")])
    _, report = report_for([sheet], [form])
    assert_error_mentions(report, "whitespace", "S1")


def test_project_filter_validates_only_samples_planned_for_requested_project(tmp_path):
    sheet = sample_sheet(
        tmp_path / "SampleSheet.csv", [["S1", PROJECT], ["S2", "GCF-2026-002"]]
    )
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("S1")])
    result, report = report_for([sheet], [form], project_ids=[PROJECT])
    assert result.ok, report
    assert report["summary"]["planned_sample_count"] == 1
    assert "S2" not in json.dumps(report["summary"]["planned_samples"])


def test_repeated_sample_across_flowcells_merges_and_warns_on_metadata_update(tmp_path):
    sheet1 = sample_sheet(tmp_path / "run1/SampleSheet.csv", [["S1", PROJECT]])
    sheet2 = sample_sheet(tmp_path / "run2/SampleSheet.csv", [["S1", PROJECT]])
    form1 = submission_form(tmp_path / "run1/submission.xlsx", [customer_row("S1", "old")])
    form2 = submission_form(tmp_path / "run2/submission.xlsx", [customer_row("S1", "new")])
    result, report = report_for([sheet1, sheet2], [form1, form2])
    assert result.ok, report
    assert report["summary"]["planned_sample_count"] == 1
    assert "s1" in diagnostics(report, "warnings")
    assert "external" in diagnostics(report, "warnings")


def test_repeated_sample_across_flowcells_can_have_different_project_ids(tmp_path):
    project2 = "GCF-2026-002"
    sheet1 = sample_sheet(tmp_path / "run1/SampleSheet.csv", [["S1", PROJECT]])
    sheet2 = sample_sheet(tmp_path / "run2/SampleSheet.csv", [["S1", project2]])
    form1 = submission_form(tmp_path / "run1/submission.xlsx", [customer_row("S1")])
    form2 = submission_form(
        tmp_path / "run2/submission.xlsx", [customer_row("S1", project=project2)]
    )
    result, report = report_for([sheet1, sheet2], [form1, form2])
    assert result.ok, report
    assert report["summary"]["planned_sample_count"] == 1
    projects = json.dumps(report["summary"]["projects"])
    assert PROJECT in projects and project2 in projects


def test_keep_batch_preserves_separate_planned_samples_from_each_flowcell(tmp_path):
    run1 = tmp_path / "260925_MACHINE_0001_FLOWCELLA"
    run2 = tmp_path / "260926_MACHINE_0002_FLOWCELLB"
    sheet1 = sample_sheet(run1 / "SampleSheet.csv", [["S1", PROJECT]])
    sheet2 = sample_sheet(run2 / "SampleSheet.csv", [["S1", PROJECT]])
    form1 = submission_form(run1 / "submission.xlsx", [customer_row("S1")])
    form2 = submission_form(run2 / "submission.xlsx", [customer_row("S1")])
    result, report = report_for([sheet1, sheet2], [form1, form2], keep_batch=True)
    assert result.ok, report
    assert report["summary"]["planned_sample_count"] == 2
    planned = json.dumps(report["summary"]["planned_samples"])
    assert "S1_FLOWCELLA" in planned
    assert "S1_FLOWCELLB" in planned


@pytest.mark.parametrize("missing_kind", ["samplesheet", "submission"])
def test_missing_file_is_reported_with_the_path(tmp_path, missing_kind):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1", PROJECT]])
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("S1")])
    missing = sheet if missing_kind == "samplesheet" else form
    missing.unlink()
    _, report = report_for([sheet], [form])
    assert_error_mentions(report, str(missing))


def test_invalid_workbook_returns_structured_error(tmp_path):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1", PROJECT]])
    form = tmp_path / "submission.xlsx"
    form.write_text("not a workbook")
    _, report = report_for([sheet], [form])
    assert_error_mentions(report, str(form))


def test_lane_specific_barcodes_are_not_identity_conflicts(tmp_path):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1", PROJECT, 1, "AAAA"], ["S1", PROJECT, 2, "CCCC"]], headers=["Sample_ID", "Sample_Project", "Lane", "index"])
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("S1")])
    result, report = report_for([sheet], [form])
    assert result.ok, report


def test_empty_csv_export_columns_and_short_optional_tail_are_accepted(tmp_path):
    sheet = tmp_path / "SampleSheet.csv"
    sheet.write_text(" [Data] ,,,\nSample_ID,Sample_Project,Description,,,\nS1," + PROJECT + "\n")
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("S1")])
    result, report = report_for([sheet], [form])
    assert result.ok, report


def test_unnamed_populated_csv_column_is_rejected(tmp_path):
    sheet = tmp_path / "SampleSheet.csv"
    sheet.write_text("[Data]\nSample_ID,Sample_Project,,\nS1," + PROJECT + ",unexpected,\n")
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("S1")])
    result, report = report_for([sheet], [form])
    assert not result.ok
    assert_error_mentions(report, "header")


def test_unknown_descriptor_dtype_is_a_header_error(tmp_path):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1", PROJECT]])
    headers = ["Unique Sample ID", "External ID", "Sample Group [dtype=typo]", "Project ID"]
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("S1")], customer_headers=headers)
    _, report = report_for([sheet], [form])
    assert_error_mentions(report, "dtype")
    assert report["errors"][0]["row"] == 15


def test_missing_metadata_has_originating_samplesheet_context(tmp_path):
    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["missing", PROJECT]])
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("other")])
    _, report = report_for([sheet], [form])
    finding = next(item for item in report["errors"] if item["code"] == "sample.missing_metadata")
    assert finding["file"] == str(sheet)
    assert finding["row"] == 5
    assert finding["column"] == "Sample_ID"


@pytest.mark.parametrize("include_demux", [True, False])
def test_parse_biosciences_uses_required_shared_demux_worksheet(tmp_path, include_demux):
    from openpyxl import load_workbook

    sheet = sample_sheet(tmp_path / "SampleSheet.csv", [["S1", PROJECT]])
    sheet.write_text(sheet.read_text().replace("[Data]", "[CustomOptions]\nLibprep,Parse Biosciences Evercode\n[Data]"))
    form = submission_form(tmp_path / "submission.xlsx", [customer_row("S1")])
    if include_demux:
        book = load_workbook(form)
        demux = book.create_sheet("Cell Multiplexing")
        demux.append(["Unique Sample ID", "Wells"])
        demux.append(["S1", "A1, A2"])
        book.save(form)
    result, report = report_for([sheet], [form])
    assert result.ok is include_demux, report
    if include_demux:
        assert result.forms[0]["demux"]["data"].iloc[0]["Sample_ID"] == "S1"
    else:
        assert_error_mentions(report, "Cell Multiplexing")
