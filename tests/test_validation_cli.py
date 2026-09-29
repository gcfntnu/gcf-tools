"""Exercise preflight ordering and output behavior through the real CLI."""

import csv
import gzip
import json
import os
from pathlib import Path
import subprocess
import sys

import pytest
import yaml
from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
PROJECT = "GCF-2026-001"
LIBPREP = "Kit SE:\n  workflow: rnaseq\n  library_layout: single\n"


def make_run(tmp_path, planned=("001",), submitted=None, discovered=("001",)):
    run = tmp_path / "input" / "260929_NB501038_0001_AEXAMPLE"
    project = run / PROJECT
    project.mkdir(parents=True)
    with (run / "SampleSheet.csv").open("w", newline="") as handle:
        handle.write("[Header]\nIEMFileVersion,4\n[Data]\n")
        writer = csv.writer(handle)
        writer.writerow(["Sample_ID", "Sample_Project"])
        writer.writerows((sample, PROJECT) for sample in planned)

    workbook = Workbook()
    customer = workbook.active
    customer.title = "Sample-Submission-Form"
    headers = ["Unique Sample ID", "External ID (optional reference sample ID)", "Sample Group", "Project ID"]
    for column, header in enumerate(headers, 1):
        customer.cell(15, column, header)
    for row, sample in enumerate(planned if submitted is None else submitted, 16):
        for column, value in enumerate([sample, "007", "control", PROJECT], 1):
            customer.cell(row, column, value)
    workbook.create_sheet("INFO (GCF-lab only)").append(["Sample_ID", "Concentration"])
    workbook.save(run / "Sample-Submission-Form.xlsx")

    stats = run / "Stats"
    stats.mkdir()
    (stats / "Stats.json").write_text(json.dumps({
        "ReadInfosForLanes": [{"LaneNumber": 1, "ReadInfos": [
            {"Number": 1, "IsIndexedRead": False, "NumCycles": 86},
            {"Number": 2, "IsIndexedRead": True, "NumCycles": 8},
        ]}],
    }))
    for sample in discovered:
        with gzip.open(project / (sample + "_R1.fastq.gz"), "wt") as handle:
            handle.write("@read\nACGT\n+\nIIII\n")
    work = tmp_path / "work"
    work.mkdir()
    return run, work


def cli(run, work, *arguments):
    # Resolve the checkout under test even when subprocess cwd is elsewhere.
    env = dict(os.environ)
    env["PYTHONPATH"] = str(ROOT) + (os.pathsep + env["PYTHONPATH"] if env.get("PYTHONPATH") else "")
    return subprocess.run(
        [sys.executable, "-m", "configmaker.configmaker", str(run),
         "--libkit", "Kit", "--organism", "Homo sapiens", "--skip-peppy", *arguments],
        cwd=work, env=env, capture_output=True, text=True, timeout=30,
    )


def install_workflow(work):
    workflow = work / "src/gcf-workflows"
    workflow.mkdir(parents=True)
    (workflow / "libprep.config").write_text(LIBPREP)


@pytest.mark.parametrize("output", ["config.yaml", "nonexistent-parent/config.yaml"])
def test_invalid_metadata_is_reported_before_any_output_or_directory_is_created(tmp_path, output):
    run, work = make_run(tmp_path, submitted=("wrong-id",))
    result = cli(run, work, "--output", output)
    assert result.returncode == 2
    assert "Metadata preflight FAILED" in result.stderr
    assert "sample.missing_metadata" in result.stderr
    assert "001" in result.stderr
    assert "Traceback" not in result.stderr
    assert not list(work.iterdir()), "Invalid metadata must not initialize output or logs"


def test_invalid_metadata_does_not_truncate_an_existing_output_file(tmp_path):
    run, work = make_run(tmp_path, submitted=("wrong-id",))
    output = work / "config.yaml"
    original = b"existing: configuration\n"
    output.write_bytes(original)
    result = cli(run, work)
    assert result.returncode == 2
    assert "sample.missing_metadata" in result.stderr
    assert output.read_bytes() == original
    assert list(work.iterdir()) == [output]


def test_valid_cli_preserves_numeric_identifiers_and_writes_discovery_summary(tmp_path):
    run, work = make_run(tmp_path)
    install_workflow(work)
    result = cli(run, work)
    assert result.returncode == 0, result.stderr
    config = yaml.safe_load((work / "config.yaml").read_text())
    assert set(config["samples"]) == {"001"}
    assert config["samples"]["001"]["Sample_ID"] == "001"
    assert config["samples"]["001"]["External_ID"] == "007"
    assert config["samples"]["001"]["Sample_Group"] == "control"
    assert config["read_geometry"] == [86]
    assert config["workflow"] == "rnaseq"
    summary = json.loads((work / "configmaker.analysis-summary.json").read_text())
    assert summary["kind"] == "fastq_discovery"
    assert summary["sample_count"] == summary["planned_sample_count"] == 1
    assert summary["missing_sample_ids"] == []
    assert summary["samples"] == [{"sample_id": "001", "flowcells": [run.name]}]
    assert "FASTQ discovery: 1 of 1 planned" in result.stdout
    assert (work / "Snakefile").is_file()
    symlinks = list((work / "data/raw/fastq").rglob("*.fastq.gz"))
    assert len(symlinks) == 1 and symlinks[0].is_symlink()
    assert symlinks[0].resolve() == run / PROJECT / "001_R1.fastq.gz"


def test_partial_fastq_discovery_reports_missing_planned_samples(tmp_path):
    run, work = make_run(tmp_path, planned=("001", "002"), submitted=("001", "002", "003"))
    install_workflow(work)
    result = cli(run, work)
    assert result.returncode == 0, result.stderr
    summary = json.loads((work / "configmaker.analysis-summary.json").read_text())
    assert summary["sample_count"] == 1
    assert summary["planned_sample_count"] == 2
    assert summary["missing_sample_ids"] == ["002"]
    assert [sample["sample_id"] for sample in summary["samples"]] == ["001"]
    assert "FASTQ discovery: 1 of 2 planned" in result.stdout
    assert "002" in (work / ".configmaker.log").read_text()
    config = yaml.safe_load((work / "config.yaml").read_text())
    assert set(config["samples"]) == {"001"}


def test_zero_discovered_fastqs_fails_without_initializing_output(tmp_path):
    run, work = make_run(tmp_path, discovered=())
    result = cli(run, work)
    assert result.returncode == 2
    assert "Metadata preflight PASSED" in result.stdout
    assert "FASTQ discovery failed" in result.stderr
    assert "no matching FASTQs" in result.stderr
    assert "Traceback" not in result.stderr
    assert not list(work.iterdir())


def test_validator_version_mismatch_fails_before_input_validation_or_output(tmp_path):
    run, work = make_run(tmp_path, submitted=("wrong-id",))
    result = cli(run, work, "--expected-validation-version", "unsupported-test-version")
    assert result.returncode == 2
    assert "validator version mismatch" in result.stderr
    assert "unsupported-test-version" in result.stderr
    assert "Metadata preflight" not in result.stderr + result.stdout
    assert "Traceback" not in result.stderr
    assert not list(work.iterdir())


def test_importing_validator_and_configmaker_creates_no_files_or_logging_handlers(tmp_path):
    script = """
import importlib
import logging
from pathlib import Path

root = logging.getLogger()
application = logging.getLogger('GCF-configmaker')
before = (list(root.handlers), list(application.handlers))
importlib.import_module('configmaker.validation')
importlib.import_module('configmaker.configmaker')
assert (list(root.handlers), list(application.handlers)) == before
assert not list(Path('.').iterdir())
"""
    env = dict(os.environ, PYTHONPATH=str(ROOT), PYTHONDONTWRITEBYTECODE="1")
    result = subprocess.run(
        [sys.executable, "-c", script], cwd=tmp_path, env=env,
        capture_output=True, text=True, timeout=30,
    )
    assert result.returncode == 0, result.stderr
    assert result.stdout == ""
    assert not list(tmp_path.iterdir())
